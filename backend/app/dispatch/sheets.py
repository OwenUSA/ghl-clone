"""Read an uploaded spreadsheet from top to bottom (2026-10-01, the owner's ask: "upload an excel,
read it from top to bottom and compare it with the information in Zuper").

`.xlsx` / `.xlsm` (openpyxl, values not formulas) or `.csv`. Every sheet, every row: the header
is the first row with at least two text cells; each row after it becomes {header: text}. Dates
become ISO dates, numbers keep their digits ("291", not "291.0"). Empty rows are skipped.
Limits keep one upload from filling the database.
"""
from __future__ import annotations

import csv
import io
from datetime import date, datetime

MAX_BYTES = 10 * 1024 * 1024
MAX_SHEETS = 20
MAX_ROWS = 20000
MAX_COLUMNS = 60
TYPES = (".xlsx", ".xlsm", ".csv")


class Unreadable(Exception):
    """A sentence for the person: why the file could not be read."""


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        if v.time() == datetime.min.time():
            return v.date().isoformat()
        return v.isoformat(" ", "minutes")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _table(name: str, raw_rows) -> dict | None:
    header: list[str] | None = None
    rows: list[dict] = []
    for n, raw in enumerate(raw_rows, start=1):
        cells = [_text(v) for v in list(raw)[:MAX_COLUMNS]]
        if header is None:
            if sum(1 for x in cells if x and not x.replace(".", "").isdigit()) >= 2:
                seen: dict[str, int] = {}
                header = []
                for i, h in enumerate(cells):
                    h = h or "Column %d" % (i + 1)
                    seen[h] = seen.get(h, 0) + 1
                    header.append(h if seen[h] == 1 else "%s (%d)" % (h, seen[h]))
            continue
        if not any(cells):
            continue
        rows.append({"row": n, "values": {header[i]: cells[i] for i in range(min(len(header),
                                                                              len(cells)))
                                          if cells[i]}})
    if header is None:
        return None
    return {"name": name, "columns": header, "rows": rows}


def parse(data: bytes, filename: str) -> list[dict]:
    """[{name, columns, rows: [{row, values}]}] — or Unreadable."""
    if len(data) > MAX_BYTES:
        raise Unreadable("The file is larger than 10 MB.")
    lower = (filename or "").lower()
    if not lower.endswith(TYPES):
        raise Unreadable("Only Excel (.xlsx, .xlsm) or .csv files can be read.")
    sheets: list[dict] = []
    if lower.endswith(".csv"):
        for enc in ("utf-8-sig", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        t = _table(filename.rsplit(".", 1)[0], csv.reader(io.StringIO(text)))
        if t:
            sheets.append(t)
    else:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001 - any damaged / non-Excel bytes: a sentence
            raise Unreadable("This Excel file could not be opened.") from None
        for ws in wb.worksheets[:MAX_SHEETS]:
            t = _table(ws.title, ws.iter_rows(values_only=True))
            if t:
                sheets.append(t)
        wb.close()
    total = 0
    for s in sheets:
        room = MAX_ROWS - total
        s["rows"] = s["rows"][:max(room, 0)]
        total += len(s["rows"])
    if not total:
        raise Unreadable("No rows with a header were found in the file.")
    return sheets


def summary(sheets: list[dict]) -> list[dict]:
    return [{"name": s["name"], "rows": len(s["rows"]), "columns": s["columns"]} for s in sheets]
