"""A week plan as an Excel workbook (2026-10-08), laid out like the office's own Schedule sheet.

  Schedule      one row per visit, a day / technician heading row above each group
  Calendar      week by week: a row per technician, a column per day, the visits in each cell
  By day        visits, driving, first start and last finish per technician per day
  Not placed    what did not fit, and why
  How it was planned   the visit lengths, the hours, the driving estimate, this plan vs a
                simple one-pass plan

A DRAFT: nothing in it is booked. The map links open Google Maps on the job's position; no
address is sent anywhere by making the file.
"""
from __future__ import annotations

import io
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from . import config as c

BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(bold=True, color="FFFFFF")
GROUP = PatternFill("solid", fgColor="D9E1F2")
BOOKED = PatternFill("solid", fgColor="EDEDED")
TENTATIVE = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
KIND = {"inspection": "Inspection", "repair": "Repair"}


def _t(iso: str | None) -> str:
    if not iso:
        return ""
    return c.local(datetime.fromisoformat(iso)).strftime("%I:%M %p").lstrip("0")


def _window(v: dict) -> str:
    return "%s-%s" % (_t(v["start"]), _t(v["end"]))


def _status(v: dict) -> str:
    if v["existing"]:
        return "Booked"
    return "Proposed - CALL FIRST (not reached yet)" if v.get("tentative") else "Proposed"


def _map(v: dict) -> str:
    if v.get("lat") is None or v.get("lng") is None:
        return ""
    return "https://www.google.com/maps/search/?api=1&query=%s,%s" % (v["lat"], v["lng"])


def _header(ws, cols: list[tuple[str, int]]) -> None:
    ws.append([h for h, _ in cols])
    for i, (_, width) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = HEAD_FONT, HEAD
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"


def _schedule(ws, plan: dict) -> None:
    _header(ws, [("Day", 12), ("Tech", 15), ("Window", 18), ("Customer", 24), ("Job #", 8),
                 ("Type", 11), ("City", 16), ("Address", 32), ("Zuper stage now", 24),
                 ("Status", 24), ("Drive before (min)", 10), ("Why it needs a visit", 40),
                 ("Customer limits", 40), ("Map", 14), ("Zuper", 14)])
    for d in plan["days"]:
        for t in d["techs"]:
            if not t["visits"]:
                continue
            ws.append(["%s  -  %s   |   %d visits, ~%d min driving, from %s" % (
                d["label"], t["tech"], t["count"], t["drive_minutes"], t["start_from"])])
            row = ws.max_row
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=15)
            ws.cell(row=row, column=1).font, ws.cell(row=row, column=1).fill = BOLD, GROUP
            for v in t["visits"]:
                ws.append([
                    datetime.fromisoformat(d["date"]).strftime("%a %m/%d"), t["tech"].split()[0],
                    _window(v), v["customer"] or "", v["job_number"] or "",
                    KIND.get(v["kind"], v["kind"]), v.get("city") or "", v.get("address") or "",
                    v.get("stage") or "", _status(v), v["drive_minutes_before"],
                    v.get("why") or "", "; ".join(v.get("limits") or []), "", ""])
                r = ws.max_row
                for col, url, text in ((14, _map(v), "Map"), (15, v.get("zuper_url"), "Open")):
                    if url:
                        cell = ws.cell(row=r, column=col, value=text)
                        cell.hyperlink, cell.font = url, Font(color="0563C1", underline="single")
                fill = BOOKED if v["existing"] else TENTATIVE if v.get("tentative") else None
                for col in range(1, 16):
                    cell = ws.cell(row=r, column=col)
                    cell.alignment, cell.border = WRAP, THIN
                    if fill:
                        cell.fill = fill


def _calendar(ws, plan: dict) -> None:
    by_day = {d["date"]: d for d in plan["days"]}
    techs = []
    for d in plan["days"]:
        for t in d["techs"]:
            if t["tech"] not in techs:
                techs.append(t["tech"])
    dates = sorted(by_day)
    if not dates:
        ws.append(["Nothing to plan."])
        return
    first = date.fromisoformat(dates[0])
    week = first - timedelta(days=first.weekday())
    last = date.fromisoformat(dates[-1])
    ws.column_dimensions["A"].width = 16
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 34
    while week <= last:
        days = [week + timedelta(days=i) for i in range(6)]
        ws.append(["Week of %s" % week.strftime("%b %d")] +
                  [x.strftime("%a %m/%d") for x in days])
        r = ws.max_row
        for col in range(1, 8):
            cell = ws.cell(row=r, column=col)
            cell.font, cell.fill = HEAD_FONT, HEAD
        for tech in techs:
            cells = [tech]
            for x in days:
                d = by_day.get(x.isoformat())
                t = next((t for t in (d or {}).get("techs", []) if t["tech"] == tech), None)
                if not t or not t["visits"]:
                    cells.append("")
                    continue
                lines = ["%d visits · ~%d min driving" % (t["count"], t["drive_minutes"])]
                for v in t["visits"]:
                    mark = "BOOKED " if v["existing"] else "CALL FIRST " if v.get("tentative") \
                        else ""
                    lines.append("%s %s#%s %s (%s) %s" % (
                        _t(v["start"]), mark, v["job_number"] or "?", v["customer"] or "",
                        v.get("city") or "", "I" if v["kind"] == "inspection" else "R"))
                cells.append("\n".join(lines))
            ws.append(cells)
            r = ws.max_row
            ws.cell(row=r, column=1).font = BOLD
            for col in range(1, 8):
                ws.cell(row=r, column=col).alignment = WRAP
                ws.cell(row=r, column=col).border = THIN
        ws.append([])
        week += timedelta(days=7)
    ws.append([("I = inspection, R = repair. BOOKED = already in Zuper. CALL FIRST = nobody "
                "has reached the customer yet.")])


def _by_day(ws, plan: dict) -> None:
    _header(ws, [("Day", 14), ("Tech", 16), ("Visits", 8), ("Of", 6), ("New", 6),
                 ("Driving (min)", 12), ("First start", 12), ("Last finish", 12)])
    for d in plan["days"]:
        for t in d["techs"]:
            new = sum(1 for v in t["visits"] if not v["existing"])
            ws.append([d["label"], t["tech"], t["count"], t["capacity"], new, t["drive_minutes"],
                       _t(t.get("first_start")), _t(t.get("last_end"))])


def _not_placed(ws, plan: dict) -> None:
    _header(ws, [("Job #", 8), ("Customer", 24), ("Type", 11), ("City", 16), ("Stage", 24),
                 ("Why not placed", 60), ("Customer limits", 40), ("Zuper", 12)])
    for u in plan["unplaced"]:
        ws.append([u["job_number"] or "", u["customer"] or "", KIND.get(u["kind"], u["kind"]),
                   u.get("city") or "", u.get("stage") or "", u["reason"],
                   "; ".join(u.get("limits") or []), ""])
        if u.get("zuper_url"):
            cell = ws.cell(row=ws.max_row, column=8, value="Open")
            cell.hyperlink, cell.font = u["zuper_url"], Font(color="0563C1", underline="single")
    if not plan["unplaced"]:
        ws.append(["Every job that needs a visit is in the plan."])


def _assumptions(ws, plan: dict) -> None:
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 90
    s, b = plan.get("summary", {}), plan.get("baseline", {})
    rows = [("Plan", "#%s, made %s" % (plan.get("id", "-"), plan.get("generated_at", ""))),
            ("Mode", "as many jobs as possible, then the least driving"
             if plan.get("mode") == "most_jobs" else "oldest waiting first"),
            ("Visits proposed", s.get("placed")), ("Not placed", s.get("not_placed")),
            ("Of them call first", s.get("tentative")),
            ("Driving, this plan (min)", s.get("drive_minutes")),
            ("Driving, a simple one-pass plan (min)", "%s for %s visits" % (
                b.get("drive_minutes"), b.get("placed"))),
            ("Arrangements tried", s.get("starts_tried")),
            ("Driving estimate", (plan.get("assumptions") or {}).get("note"))]
    for name, m in ((plan.get("assumptions") or {}).get("minutes") or {}).items():
        rows.append(("Visit length - %s" % name, "inspection %s min, repair %s min (%s)" % (
            m.get("inspection"), m.get("repair"), m.get("source"))))
    for name, h in ((plan.get("assumptions") or {}).get("hours") or {}).items():
        rows.append(("Hours - %s" % name, "%s-%s, up to %s visits a day" % (
            h.get("start"), h.get("end"), h.get("max"))))
    rows.append(("Draft", ("Nothing here is booked. The office books each visit in Zuper "
                           "after agreeing it with the customer.")))
    for k, v in rows:
        ws.append([k, "" if v is None else v])
        ws.cell(row=ws.max_row, column=1).font = BOLD


def workbook(plan: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    _schedule(ws, plan)
    _calendar(wb.create_sheet("Calendar"), plan)
    _by_day(wb.create_sheet("By day"), plan)
    _not_placed(wb.create_sheet("Not placed"), plan)
    _assumptions(wb.create_sheet("How it was planned"), plan)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
