"""/api/dispatch/chats — the Dispatch assistant's conversations, SAVED (2026-10-01).

The owner's ask: keep every chat and everything the assistant did in it, so it can be reviewed
and improved later, and list the chats beside the conversation.

  * Every question, every answer, every tool the assistant used (what it looked up and what it
    got back), the model and tokens, and any error — `dispatch_chats` /
    `dispatch_chat_messages`. The server keeps the history: the browser sends only the new
    question, so a conversation cannot be edited after the fact.
  * A person's thumbs up / down and note on an answer — the improvement signal.
  * "Delete" only hides a chat from its owner's list (`archived`); it stays for review.

Who: the Dispatch audience (ADMIN + unrestricted DISPATCHER). A person sees and writes only
their own chats (another's answers 404, as an id that does not exist would); an ADMIN can also
read everyone's (`scope=all`), read-only. The assistant's tools see only the boards the asking
user may see.
"""
from __future__ import annotations

import io
from datetime import UTC, datetime

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .. import auth
from ..db import get_db
from ..models import (
    DispatchChat,
    DispatchChatFile,
    DispatchChatMessage,
    DispatchJob,
    DispatchSuggestion,
    Role,
    User,
)
from . import ai, compare, sheets
from . import config as c
from .api import VIEW, _hidden, _suggestion

router = APIRouter(prefix="/api/dispatch/chats", tags=["dispatch"])

TITLE_MAX = 80
HISTORY_MAX = 30           # messages sent to the model with a new question


def _iso(dt) -> str | None:
    dt = c.aware(dt)
    return dt.isoformat() if dt else None


def _summary(db: Session, ch: DispatchChat, count: int | None = None,
             name: str | None = None) -> dict:
    if count is None:
        count = db.scalar(select(func.count()).select_from(DispatchChatMessage).where(
            DispatchChatMessage.chat_id == ch.id)) or 0
    if name is None and ch.user_id is not None:
        u = db.get(User, ch.user_id)
        name = u.name if u else None
    return {"id": ch.id, "title": ch.title, "user_id": ch.user_id, "user_name": name,
            "created_at": _iso(ch.created_at), "updated_at": _iso(ch.updated_at),
            "message_count": count}


def _message(db: Session, m: DispatchChatMessage) -> dict:
    sugs = db.scalars(select(DispatchSuggestion).where(
        DispatchSuggestion.id.in_(m.suggestion_ids))).all() if m.suggestion_ids else []
    return {"id": m.id, "role": m.role, "content": m.content, "steps": m.steps or [],
            "error": m.error, "model": m.model,
            "feedback": {"rating": m.rating, "note": m.feedback_note} if m.rating else None,
            "suggestions": [_suggestion(s) for s in sugs], "created_at": _iso(m.created_at),
            "files": [_file(f) for f in db.scalars(select(DispatchChatFile).where(
                DispatchChatFile.message_id == m.id).order_by(DispatchChatFile.id))]
            if m.role == "user" else [],
            "downloads": m.downloads or []}


def _file(f: DispatchChatFile) -> dict:
    return {"id": f.id, "filename": f.filename, "total_rows": f.total_rows,
            "sheets": [{"name": x["name"], "rows": len(x["rows"]), "columns": x["columns"]}
                       for x in f.sheets]}


def _own(db: Session, principal: auth.Principal, chat_id: int, *,
         write: bool) -> DispatchChat:
    """The chat, if this person may see it (theirs; or any, read-only, for an ADMIN)."""
    ch = db.get(DispatchChat, chat_id)
    mine = ch is not None and ch.user_id == principal.user_id
    if ch is None or (not mine and (write or principal.role is not Role.ADMIN)) or \
            (ch.archived and not (principal.role is Role.ADMIN and not write)):
        raise HTTPException(404, "no such chat")
    return ch


@router.get("")
def dispatch_list_chats(scope: str = "mine", principal: auth.Principal = VIEW,
                        db: Session = Depends(get_db)):
    if scope not in ("mine", "all"):
        raise HTTPException(400, "scope is mine or all")
    if scope == "all" and principal.role is not Role.ADMIN:
        raise HTTPException(403, "only an admin can read everyone's chats")
    counts = select(DispatchChatMessage.chat_id, func.count().label("n")) \
        .group_by(DispatchChatMessage.chat_id).subquery()
    stmt = (select(DispatchChat, counts.c.n, User.name)
            .outerjoin(counts, counts.c.chat_id == DispatchChat.id)
            .outerjoin(User, User.id == DispatchChat.user_id)
            .order_by(DispatchChat.updated_at.desc()).limit(300))
    if scope == "mine":
        stmt = stmt.where(DispatchChat.user_id == principal.user_id,
                          DispatchChat.archived.is_(False))
    return {"chats": [_summary(db, ch, n or 0, name) for ch, n, name in db.execute(stmt)]}


@router.post("/files")
async def dispatch_upload_chat_file(file: UploadFile = File(...),
                                   principal: auth.Principal = VIEW,
                                   db: Session = Depends(get_db)):
    """A spreadsheet for the assistant: read now, from top to bottom, kept with the chat once a
    message carries it. Only the rows are stored, not the file."""
    data = await file.read(sheets.MAX_BYTES + 1)
    if len(data) > sheets.MAX_BYTES:
        raise HTTPException(413, "The file is larger than 10 MB.")
    if not (file.filename or "").lower().endswith(sheets.TYPES):
        raise HTTPException(415, "Only Excel (.xlsx, .xlsm) or .csv files can be read.")
    try:
        parsed = sheets.parse(data, file.filename or "upload.xlsx")
    except sheets.Unreadable as e:
        raise HTTPException(422, str(e)) from None
    f = DispatchChatFile(user_id=principal.user_id, filename=(file.filename or "upload")[:255],
                         size=len(data), total_rows=sum(len(x["rows"]) for x in parsed),
                         sheets=parsed, created_at=datetime.now(UTC))
    db.add(f)
    db.commit()
    return _file(f)


def _own_file(db: Session, principal: auth.Principal, file_id: int) -> DispatchChatFile:
    f = db.get(DispatchChatFile, file_id)
    if f is None or (f.user_id != principal.user_id and principal.role is not Role.ADMIN):
        raise HTTPException(404, "no such file")
    return f


@router.get("/files/{file_id}/comparison.xlsx")
def dispatch_file_comparison(file_id: int, principal: auth.Principal = VIEW,
                             db: Session = Depends(get_db)):
    """The row-by-row comparison of an uploaded spreadsheet with Zuper, as an Excel workbook.
    Built fresh from Zuper's current copy each time."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    f = _own_file(db, principal, file_id)
    result = compare.compare(f.sheets, db.scalars(select(DispatchJob)).all(),
                             hidden=_hidden(db, principal), now=datetime.now(UTC))
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"
    head = ["Result", "Sheet", "Row", "Customer (file)", "Status (file)", "Matched by",
            "Zuper job #", "Board", "Zuper stage now", "In stage since", "Next visit",
            "Assigned", "Differences", "Other jobs"]
    ws.append(head)
    fills = {"different": "FFE699", "not in Zuper": "F8CBAD", "same": "E2EFDA"}
    for x in result["rows"]:
        ws.append([x["result"], x["sheet"], x["row"], x["customer"], x["file_status"],
                   x["matched_by"], x["job_number"], x["board"], x["zuper_stage"],
                   x["stage_since"], x["visit"], x["assigned"], " ".join(x["differences"]),
                   ", ".join(x["other_jobs"])])
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=fills.get(x["result"],
                                                                             "FFFFFF"))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCDEFGHIJKLMN", (14, 18, 6, 26, 24, 12, 10, 22, 24, 20, 20, 18,
                                             70, 30), strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    sm = wb.create_sheet("Summary")
    sm.append(["Compared", f.filename])
    sm.append(["Made", datetime.now(UTC).astimezone(c.TZ).strftime("%b %d, %Y %I:%M %p")])
    for k, v in result["counts"].items():
        sm.append([k.replace("_", " "), v])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f.filename.rsplit(".", 1)[0][:80] + " - compared with Zuper.xlsx"
    return StreamingResponse(buf, media_type=(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        headers={"Content-Disposition": 'attachment; filename="%s"' % name.replace('"', "")})


@router.get("/{chat_id}")
def dispatch_get_chat(chat_id: int, principal: auth.Principal = VIEW,
                      db: Session = Depends(get_db)):
    ch = _own(db, principal, chat_id, write=False)
    msgs = db.scalars(select(DispatchChatMessage).where(DispatchChatMessage.chat_id == ch.id)
                      .order_by(DispatchChatMessage.id)).all()
    return {"chat": _summary(db, ch, len(msgs)),
            "can_write": ch.user_id == principal.user_id and not ch.archived,
            "messages": [_message(db, m) for m in msgs]}


class AskIn(BaseModel):
    chat_id: int | None = None
    content: str = Field(min_length=1, max_length=4000)
    file_ids: list[int] = Field(default_factory=list, max_length=3)


@router.post("/messages")
def dispatch_chat_message(body: AskIn, principal: auth.Principal = VIEW,
                          db: Session = Depends(get_db)):
    """Ask the assistant. A new chat when `chat_id` is null. The history comes from the
    server; the question, the answer and every tool step are saved — an unavailable AI too."""
    text = body.content.strip()
    if not text:
        raise HTTPException(422, "the question is empty")
    now = datetime.now(UTC)
    if body.chat_id is None:
        ch = DispatchChat(user_id=principal.user_id, title=text[:TITLE_MAX], created_at=now,
                          updated_at=now)
        db.add(ch)
        db.flush()
    else:
        ch = _own(db, principal, body.chat_id, write=True)
    history = db.scalars(select(DispatchChatMessage).where(
        DispatchChatMessage.chat_id == ch.id, DispatchChatMessage.error.is_(False))
        .order_by(DispatchChatMessage.id.desc()).limit(HISTORY_MAX)).all()[::-1]
    attached = []
    for fid in dict.fromkeys(body.file_ids):
        f = db.get(DispatchChatFile, fid)
        if f is None or f.user_id != principal.user_id or (
                f.chat_id is not None and f.chat_id != ch.id):
            raise HTTPException(404, "no such file")
        attached.append(f)
    asked = DispatchChatMessage(chat_id=ch.id, role="user", content=text, created_at=now)
    db.add(asked)
    db.flush()
    for f in attached:
        f.chat_id, f.message_id = ch.id, asked.id
    db.commit()
    files = {f.id: f for f in db.scalars(select(DispatchChatFile).where(
        DispatchChatFile.chat_id == ch.id))}
    question = text
    if attached:
        question += "\n\n" + "\n".join(
            "[Attached file #%d “%s”: %s]" % (f.id, f.filename, "; ".join(
                "sheet “%s” %d rows, columns %s" % (x["name"], len(x["rows"]),
                                                    ", ".join(x["columns"][:15]))
                for x in f.sheets)) for f in attached)
    out = ai.chat(db, [{"role": m.role, "content": m.content} for m in history] +
                  [{"role": "user", "content": question}],
                  user_id=principal.user_id, hidden=_hidden(db, principal), now=now,
                  files=files)
    answer = DispatchChatMessage(
        chat_id=ch.id, role="assistant", content=out["reply"] or "(no answer)",
        steps=out["steps"] or None, suggestion_ids=out["suggestions"] or None,
        error=bool(out["error"]), model=out["model"], input_tokens=out["input_tokens"],
        downloads=out.get("downloads") or None,
        output_tokens=out["output_tokens"], created_at=datetime.now(UTC))
    db.add(answer)
    ch.updated_at = datetime.now(UTC)
    db.commit()
    return {"chat": _summary(db, ch), "user_message": _message(db, asked),
            "assistant_message": _message(db, answer)}


class RenameIn(BaseModel):
    title: str = Field(min_length=1, max_length=TITLE_MAX)


@router.patch("/{chat_id}")
def dispatch_rename_chat(chat_id: int, body: RenameIn, principal: auth.Principal = VIEW,
                         db: Session = Depends(get_db)):
    ch = _own(db, principal, chat_id, write=True)
    ch.title = body.title.strip() or ch.title
    db.commit()
    return _summary(db, ch)


@router.delete("/{chat_id}")
def dispatch_delete_chat(chat_id: int, principal: auth.Principal = VIEW,
                         db: Session = Depends(get_db)):
    """Hides the chat from its owner's list. It is KEPT (archived) for review."""
    ch = _own(db, principal, chat_id, write=True)
    ch.archived = True
    db.commit()
    return {"id": ch.id, "archived": True}


class FeedbackIn(BaseModel):
    rating: int = Field(ge=-1, le=1)
    note: str | None = Field(default=None, max_length=2000)


@router.post("/{chat_id}/messages/{message_id}/feedback")
def dispatch_chat_feedback(chat_id: int, message_id: int, body: FeedbackIn,
                           principal: auth.Principal = VIEW, db: Session = Depends(get_db)):
    """Thumbs up (1) / down (-1) on one of the assistant's answers, with an optional note."""
    if body.rating == 0:
        raise HTTPException(422, "rating is 1 or -1")
    ch = _own(db, principal, chat_id, write=True)
    m = db.get(DispatchChatMessage, message_id)
    if m is None or m.chat_id != ch.id or m.role != "assistant":
        raise HTTPException(404, "no such answer")
    m.rating = body.rating
    m.feedback_note = (body.note or "").strip() or None
    m.feedback_at = datetime.now(UTC)
    db.commit()
    return _message(db, m)
