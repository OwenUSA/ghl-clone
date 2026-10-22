"""The report as an Excel workbook with native charts (openpyxl): Summary (each KPI, this
period vs the previous one), one sheet per table, a Trend sheet, and bar / line charts. Styled
like the nightly KPI Excel (`kpi_report._sheet`). Built on demand; nothing is stored."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference

from .. import kpi_report as kr
from . import metrics

TABLE_TITLES = {"technicians": "Technicians", "callbacks": "Callbacks",
                "still_leaking": "Still leaking", "complaints": "Complaints",
                "callback_requests": "Asked to come back", "upsells": "Upsells",
                "ratings": "Ratings"}


def workbook(rep: dict, prev: dict | None, *, label: str) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    tiles, before = rep["tiles"], (prev or {}).get("tiles", {})
    rows = []
    for k, v in tiles.items():
        p = before.get(k)
        rows.append([metrics.TILE_LABELS[k], v, p, (v - p) if isinstance(p, (int, float))
                     and isinstance(v, (int, float)) else None])
    ws = kr._sheet(wb, "Summary", ["KPI (%s)" % label, "This period", "Previous", "Change"],
                   rows, {"KPI (%s)" % label: 34})
    ws.cell(row=len(rows) + 3, column=1).value = "Period: %s to %s" % (rep["first"],
                                                                        rep["last"])
    for name, title in TABLE_TITLES.items():
        data = rep["tables"].get(name) or []
        cols = list(data[0].keys()) if data else ["(nothing in this period)"]
        sheet = kr._sheet(wb, title, [c.replace("_", " ").capitalize() for c in cols],
                          [[r.get(c) for c in cols] for r in data],
                          {c.replace("_", " ").capitalize(): 22 for c in cols})
        if name == "technicians" and data:
            chart = BarChart()
            chart.title = "Callbacks and still leaking per technician"
            chart.y_axis.title = "Jobs"
            ci, si = cols.index("callbacks") + 1, cols.index("still_leaking") + 1
            for col in (ci, si):
                chart.add_data(Reference(sheet, min_col=col, min_row=1, max_row=len(data) + 1),
                               titles_from_data=True)
            chart.set_categories(Reference(sheet, min_col=1, min_row=2,
                                           max_row=len(data) + 1))
            chart.height, chart.width = 8, 18
            sheet.add_chart(chart, "J2")
    series = rep["series"]
    keys = sorted({k for s in series.values() for k in s})
    names = [n for n in ("leads", "inspections", "repairs", "callbacks", "still_leaking")
             if n in series]
    if keys and names:
        trend = kr._sheet(wb, "Trend", ["Period"] + [n.replace("_", " ").capitalize()
                                                       for n in names],
                          [[k] + [series[n].get(k, 0) for n in names] for k in keys])
        line = LineChart()
        line.title = "Over the period"
        for i in range(len(names)):
            line.add_data(Reference(trend, min_col=i + 2, min_row=1, max_row=len(keys) + 1),
                          titles_from_data=True)
        line.set_categories(Reference(trend, min_col=1, min_row=2, max_row=len(keys) + 1))
        line.height, line.width = 9, 20
        trend.add_chart(line, "H2")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def result_workbook(result: dict) -> bytes:
    """One "Ask" result as a workbook (with its chart when it has one)."""
    wb = Workbook()
    wb.remove(wb.active)
    cols = [str(c) for c in result.get("columns") or []]
    rows = result.get("rows") or []
    ws = kr._sheet(wb, "Result", cols, rows, dict.fromkeys(cols, 22))
    ws.cell(row=len(rows) + 3, column=1).value = result.get("title") or ""
    if rows and len(cols) >= 2 and result.get("chart") in ("bar", "line"):
        num = next((i + 1 for i, _ in enumerate(cols)
                    if i and all(isinstance(r[i], (int, float)) or r[i] is None for r in rows)),
                   None)
        if num:
            ch = BarChart() if result["chart"] == "bar" else LineChart()
            ch.title = result.get("title") or ""
            ch.add_data(Reference(ws, min_col=num, min_row=1, max_row=len(rows) + 1),
                        titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(rows) + 1))
            ws.add_chart(ch, "H2")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
