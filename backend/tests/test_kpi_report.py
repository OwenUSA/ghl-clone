"""The nightly KPI Excel (2026-09-30), on made-up data only.

Pins the owner's definitions — an upsell is a job closed for more than the AHS invoice, the
difference is what the customer paid — the address matching that found the real 31 upsells,
that the file carries no customer details, and that a day's file is never overwritten.
"""
import csv
from datetime import UTC, datetime, timedelta
from pathlib import Path

from app import kpi_report
from app.db import SessionLocal
from app.models import ZuperRecordVersion, ZuperStatusHistory
from openpyxl import load_workbook

WORKIZ_HEADERS = ["Job #", "Job name", "Client", "Tags", "Type", "Job Created", "Scheduled",
                  "End", "Phone", "Email", "Status", "Tech", "Created by", "Address", "City",
                  "State", "Zip code", "Total", "Source", "Lead Created Date", "Job origin"]
AHS_HEADERS = ["Source Page", "Dispatch ID", "Covered Address", "Street", "City", "State",
               "ZIP", "Vendor ID", "Invoice Submitted", "Invoice Number", "Item",
               "Item Amount", "Status", "Invoices on Dispatch", "Dispatch Total"]
INVOICE_HEADERS = ["Invoice NO.", "Invoice Name", "Client", "Email Address", "Created",
                   "Subtotal", "Discount", "Tax", "Total Amount", "Amount Due", "Status", "Job",
                   "Job name"]
SECRET_NAME = "Zelda Customerperson"


def write_csv(path: Path, headers: list[str], rows: list[dict]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=headers)
        w.writeheader()
        for r in rows:
            w.writerow({h: r.get(h, "") for h in headers})


def workiz(number, address, created, total, *, zip_code="34205", source="AHS", status="Done",
           tech="Antonio Brown", tags=""):
    return {"Job #": number, "Client": SECRET_NAME, "Phone": "+15555550123",
            "Email": "zelda@example.test", "Address": address, "Zip code": zip_code,
            "Job Created": created, "Total": total, "Source": source, "Status": status,
            "Tech": tech, "Type": "Active Leak Repair", "Tags": tags}


def dispatch(did, street, submitted, amount, zip_code="34205"):
    return {"Dispatch ID": did, "Street": street, "ZIP": zip_code, "Covered Address": street,
            "Invoice Submitted": submitted, "Item Amount": amount, "Dispatch Total": amount,
            "Status": "Approved", "Invoice Number": "1" + did[-2:]}


def inputs(tmp_path: Path) -> Path:
    d = tmp_path / "input"
    d.mkdir()
    write_csv(d / "workiz_jobs_16.csv", WORKIZ_HEADERS, [
        workiz("AAA111", "100 Palm St", "Mon Jan 05, 2026 09:00 am", "2600.00"),     # upsell
        workiz("BBB222", "200 Oak Ave", "Mon Jan 05, 2026 09:00 am", "1100.00"),     # AHS only
        # House number 34205 would be a ZIP if read from anywhere but the END of the address.
        workiz("CCC333", "34205 Pine Rd, Bradenton, FL 34211", "Tue Feb 03, 2026 10:00 am",
               "1500.00", zip_code=""),
        workiz("DDD444", "9 Retail Way", "Tue Feb 03, 2026 10:00 am", "900.00",
               source="Google", status="Canceled", tags="Callback"),
    ])
    # A later export wins: BBB222's total is corrected here.
    write_csv(d / "workiz_jobs_19.csv", WORKIZ_HEADERS, [
        workiz("BBB222", "200 Oak Ave", "Mon Jan 05, 2026 09:00 am", "1000.00")])
    write_csv(d / "submitted_invoices_vendor_677998.csv", AHS_HEADERS, [
        dispatch("11111101", "100 Palm St", "01/12/2026", "1100.00"),
        dispatch("11111102", "200 Oak Ave", "01/13/2026", "1000.00"),
        dispatch("11111103", "34205 Pine Rd", "02/10/2026", "1100.00", zip_code="34211"),
    ])
    # AAA111's invoice (not its job total) is what it closed for.
    write_csv(d / "workiz-invoices.csv", INVOICE_HEADERS, [
        {"Invoice NO.": "AAA111", "Job": "AAA111", "Client": SECRET_NAME,
         "Total Amount": "3000.00", "Amount Due": "0.00", "Created": "Mon Jan 12, 2026",
         "Status": "Paid - sent on Mon Jan 12, 2026 01:00 pm"}])
    return d


def test_the_upsell_rule_and_the_address_match(tmp_path):
    d = inputs(tmp_path)
    jobs = kpi_report.load_workiz_jobs(d)
    assert jobs["BBB222"].total == 1000.0                       # the later export won
    matches = kpi_report.match_dispatches(kpi_report.load_ahs(d), jobs)
    rows = {r.job: r for r in kpi_report.upsells(matches, kpi_report.load_workiz_invoices(d))}
    assert rows["AAA111"].outcome == "UPSELL" and rows["AAA111"].difference == 1900.0
    assert rows["AAA111"].closed_from == "Workiz invoice"
    assert rows["BBB222"].outcome == "AHS only"
    assert rows["CCC333"].outcome == "UPSELL" and rows["CCC333"].method == "address + ZIP"
    assert "DDD444" not in rows                                  # Retail: no AHS dispatch


def test_one_dispatch_per_job_nearest_in_time():
    job = kpi_report.WorkizJob("J1", True, "Done", "", "", "", 1100.0,
                               datetime(2026, 3, 1), "5", "ELM", "34205")
    def at(did, when):
        return kpi_report.Dispatch(did, 1100.0, when, "Approved", "5", "ELM", "34205")
    early, late, before = (at("D1", datetime(2026, 3, 4)), at("D2", datetime(2026, 3, 30)),
                           at("D3", datetime(2026, 2, 1)))
    got = kpi_report.match_dispatches({"D1": early, "D2": late, "D3": before}, {"J1": job})
    methods = {m.dispatch.dispatch_id: m.method for m in got}
    assert methods["D1"] == "address + ZIP"                      # the nearest takes the job
    assert methods["D2"].endswith("(shared job)")
    assert "D3" not in methods                                   # submitted before the job


def seed_zuper(now: datetime) -> None:
    def job(uid, number, board, status, stype, *, checklist=None, tags=None, job_type=None):
        return {"job_uid": uid, "work_order_number": number, "is_deleted": False,
                "job_category": {"category_uid": "cat-" + board, "category_name": board},
                "current_job_status": {"status_name": status, "status_type": stype},
                "job_title": SECRET_NAME + " - roof", "customer": {"customer_first_name": "Zelda"},
                "custom_fields": [{"label": "Technician", "value": "Antonio Brown"},
                                  {"label": "Job Type", "value": job_type or "Roof Repair"}],
                "job_tags": tags or [], "created_at": "2026-09-27T12:00:00Z",
                "job_status": [{"status_name": status, "checklist": checklist or []}]}
    records = [
        job("j1", 701, "AHS - Repair & Review", "Invoice Submitted to AHS", "OTHER", checklist=[
            {"question": "Customer chose", "answer": "Better"},
            {"question": "Customer paid ($)", "answer": "850"}]),
        job("j2", 702, "AHS - Repair & Review", "Paid", "PAID"),
        job("j3", 703, "Retail", "Cancelled", "CANCELED", checklist=[
            {"question": "Cancel reason", "answer": "Price"}]),
        job("j4", 704, "Retail", "New Lead", "NEW", job_type="Callback/Warranty"),
        job("j9", 709, "Some other board", "New", "NEW"),
    ]
    rr = "cat-AHS - Repair & Review"
    with SessionLocal() as s:
        for r in records:
            s.add(ZuperRecordVersion(module="job", zuper_uid=r["job_uid"], content_hash="h",
                                     record=r))
        t0 = now - timedelta(hours=30)
        s.add_all([
            ZuperStatusHistory(history_uid="h1", job_uid="j1", category_uid=rr,
                               status_name="Repair Complete", changed_at=t0, done_by_name="Luis"),
            ZuperStatusHistory(history_uid="h2", job_uid="j1", category_uid=rr,
                               status_name="Invoice Submitted to AHS",
                               changed_at=t0 + timedelta(hours=6), done_by_name="Luis"),
            # Before 2026-09-26: the load's moves are not timing data.
            ZuperStatusHistory(history_uid="h0", job_uid="j2", category_uid=rr,
                               status_name="Paid", changed_at=datetime(2026, 9, 17, tzinfo=UTC)),
        ])
        s.commit()


def cells(ws) -> list[list]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


def test_the_workbook_answers_the_owner_s_questions_without_customer_details(tmp_path, db):
    now = datetime(2026, 10, 1, 7, 0, tzinfo=UTC)
    seed_zuper(now)
    with SessionLocal() as s:
        path, info = kpi_report.write(s, inputs(tmp_path), tmp_path / "out", now)
    assert info["upsells"] == 2 and info["upsell_customer_paid"] == 2300.0
    wb = load_workbook(path)
    assert wb.sheetnames[0] == "Read me"
    for name in ("Summary", "AHS upsells", "Upsells by month", "Upsells by tech",
                 "Workiz by month", "Jobs now", "Time in column", "Moves"):
        assert name in wb.sheetnames

    summary = {(r[0], r[1]): r[2] for r in cells(wb["Summary"])}
    assert summary[("AHS - Repair & Review", "closed (Paid / Completed)")] == 1
    assert summary[("Retail", "cancelled")] == 1
    assert summary[("Retail", "callbacks (re-work / warranty)")] == 1
    assert summary[("AHS", "UPSELLS (closed > AHS invoice)")] == 2
    assert summary[("AHS - Repair & Review",
                    "upsells recorded in Zuper (\"Customer chose\" not AHS only)")] == 1

    jobs_now = {r[0]: r for r in cells(wb["Jobs now"])[1:]}
    assert set(jobs_now) == {701, 702, 703, 704}                 # other boards left out
    head = cells(wb["Jobs now"])[0]
    assert jobs_now[701][head.index("Customer chose")] == "Better"
    assert jobs_now[703][head.index("Cancel reason")] == "Price"

    timing = {(r[0], r[1]): r for r in cells(wb["Time in column"])[1:]}
    left = timing[("AHS - Repair & Review", "Repair Complete")]
    assert left[2] == 1 and left[3] == 6.0                        # left after 6 hours
    assert ("AHS - Repair & Review", "Paid") not in timing        # the load's move is ignored

    text = " ".join(str(v) for ws in wb.worksheets for row in cells(ws) for v in row)
    for secret in (SECRET_NAME, "Zelda", "+15555550123", "zelda@example.test", "Palm St"):
        assert secret not in text


def test_a_corrected_job_field_wins_over_the_checklist_answer():
    job = {"job_status": [{"checklist": [{"question": "Customer paid ($)", "answer": "85"},
                                         {"question": "Customer chose", "answer": "Better"}]}],
           "custom_fields": [{"label": "Customer paid ($)", "value": "850"},
                             {"label": "Customer chose", "value": ""}]}
    got = kpi_report.answers(job)
    assert got["Customer paid ($)"] == "850"          # corrected on the job page
    assert got["Customer chose"] == "Better"          # empty field: the checklist answer stands


def job_with_items(tech="Antonio Brown", day="2026-08-10T14:00:00Z", items=()):
    return {"scheduled_start_time": day, "custom_fields": [{"label": "Technician", "value": tech}],
            "products": [{"product_id": c, "price": p, "quantity": q, "total": p * q}
                         for c, p, q in items]}


def test_expected_commission_follows_the_rule_per_line():
    job = job_with_items(items=[("AHS-LEAK-REPAIR", 1100, 1), ("UPG-UNKNOWN", 900, 1)])
    assert kpi_report.expected_commission(job, kpi_report.DEFAULT_RULES) == 350 + 450
    trip = job_with_items(items=[("AHS-TRIP", 100, 1)])
    assert kpi_report.expected_commission(trip, kpi_report.DEFAULT_RULES) == 50
    two_leaks = job_with_items(items=[("AHS-LEAK-REPAIR", 1100, 2)])
    assert kpi_report.expected_commission(two_leaks, kpi_report.DEFAULT_RULES) == 700
    # No rule for this technician: no expectation, rather than a wrong one.
    nico = job_with_items(tech="Nico")
    assert kpi_report.expected_commission(nico, kpi_report.DEFAULT_RULES) is None


def test_a_new_rule_on_the_server_changes_only_jobs_from_its_date(tmp_path):
    (tmp_path / "commission-rules.csv").write_text(
        "technician,item,kind,value,from\n"
        "Antonio Brown,AHS-LEAK-REPAIR,flat,350,2026-01-01\n"
        "Antonio Brown,AHS-LEAK-REPAIR,flat,400,2026-11-01\n", encoding="utf-8")
    rules, source = kpi_report.load_rules(tmp_path)
    assert source == "commission-rules.csv"
    old = job_with_items(day="2026-10-15T12:00:00Z", items=[("AHS-LEAK-REPAIR", 1100, 1)])
    new = job_with_items(day="2026-11-03T12:00:00Z", items=[("AHS-LEAK-REPAIR", 1100, 1)])
    assert kpi_report.expected_commission(old, rules) == 350
    assert kpi_report.expected_commission(new, rules) == 400


def test_income_per_job_and_commissions_by_week(tmp_path, db):
    now = datetime(2026, 10, 2, 7, 0, tzinfo=UTC)
    job = {"job_uid": "j5", "work_order_number": 805, "is_deleted": False, "job_total": 2000,
           "job_category": {"category_uid": "c", "category_name": "AHS - Repair & Review"},
           "current_job_status": {"status_name": "Paid", "status_type": "PAID"},
           "custom_fields": [{"label": "Technician", "value": "Antonio Brown"}],
           "scheduled_start_time": "2026-09-28T13:00:00Z", "created_at": "2026-09-17T12:00:00Z",
           "products": [
               {"product_id": "AHS-LEAK-REPAIR", "price": 1100, "quantity": 1, "total": 1100},
               {"product_id": "UPG-BETTER", "price": 900, "quantity": 1, "total": 900}]}
    comm = {"commission_uid": "c1", "job_uid": "j5", "job": {"work_order_number": 805},
            "commission_amount": 350, "commission_date": "2026-09-29", "payout_status": "UNPAID",
            "assigned_to": {"first_name": "Antonio", "last_name": "Brown"}, "is_deleted": False}
    with SessionLocal() as s:
        s.add(ZuperRecordVersion(module="job", zuper_uid="j5", content_hash="h", record=job))
        s.add(ZuperRecordVersion(module="commission", zuper_uid="c1", content_hash="h",
                                 record=comm))
        s.commit()
        path, _ = kpi_report.write(s, tmp_path / "none", tmp_path / "out", now)
    wb = load_workbook(path)
    head, *rows = cells(wb["Income per job"])
    row = dict(zip(head, rows[0], strict=True))
    assert row["AHS paid"] == 1100 and row["Customer paid (upgrade)"] == 900
    assert row["Upgrade tier"] == "Better"
    assert row["Commission in Zuper"] == 350 and row["Expected commission (rule)"] == 800
    assert row["Difference (Zuper - rule)"] == -450          # the upgrade commission is missing
    weekly = cells(wb["Commissions by week"])[1]
    assert weekly == ["Antonio Brown", "2026-09-28", 1, 350, 0, 350]
    summary = {(r[0], r[1]): r[2] for r in cells(wb["Summary"])}
    assert summary[("AHS", "customers paid on upgrades ($)")] == 900
    assert summary[("All", "still owed ($)")] == 350


def test_a_copied_pay_sheet_tab_is_counted_once_under_its_real_period(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    rows = [(datetime(2026, 8, 3), "AHS", "x", "y", 1100.0, 350.0),
            (datetime(2026, 8, 11), "AHS", "x", "y", 1125.0, 350.0)]
    for title in ("07152026 to 08012026", "08012026 to 08152026"):     # the copy comes FIRST
        ws = wb.create_sheet(title)
        ws.append(["Date", "Company", "Name", "Address", "Price", "Antonio"])
        for r in rows:
            ws.append(list(r))
    wb.save(tmp_path / "antonio-commissions.xlsx")
    got = {r[0]: r for r in kpi_report.pay_sheet_periods(tmp_path)}
    assert got["08012026 to 08152026"][5] == 700                 # the real period is counted
    assert got["07152026 to 08012026"][6].startswith("copy of tab 08012026")


def test_a_day_s_file_is_never_overwritten(tmp_path, db):
    now = datetime(2026, 10, 1, 7, 0, tzinfo=UTC)
    with SessionLocal() as s:
        first, _ = kpi_report.write(s, tmp_path / "none", tmp_path / "out", now)
        second, _ = kpi_report.write(s, tmp_path / "none", tmp_path / "out",
                                     now + timedelta(minutes=5))
    assert first != second and first.exists() and second.exists()


def test_the_nightly_hook_never_raises(monkeypatch, db):
    def boom(*a, **k):
        raise RuntimeError("disk full")
    monkeypatch.setattr(kpi_report, "write", boom)
    with SessionLocal() as s:
        kpi_report.nightly(s)
