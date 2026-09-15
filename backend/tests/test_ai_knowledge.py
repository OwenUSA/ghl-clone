"""Knowledge bases: PDF and Word files become searchable text, retrieval never leaves the
agent's attached knowledge bases, gaps are recorded once per question and resolving one
writes the FAQ."""
import base64
import io

from app.ai import knowledge
from app.db import SessionLocal
from app.models import AiKbChunk, AiKbItem, AiKnowledgeGap
from sqlalchemy import select
from tests.ai_support import anthropic_message, count, drain, make_agent, text, tool_use


def _pdf_bytes(sentence: str) -> bytes:
    """A real one-page PDF with `sentence` on it, written by hand (no PDF writer needed)."""
    content = ("BT /F1 12 Tf 72 720 Td (%s) Tj ET" % sentence).encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
         b"/Resources << /Font << /F1 5 0 R >> >> >>"),
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
              % (len(objects) + 1, xref))
    return out.getvalue()


def _docx_bytes(*paragraphs: str) -> bytes:
    import docx
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Shingle warranty"
    table.rows[0].cells[1].text = "25 years"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_pdf_and_docx_text_is_extracted_and_searchable(world):
    admin = world.client("admin")
    kb = admin.post("/api/ai/knowledge-bases", json={"name": "Services"}).json()
    pdf = _pdf_bytes("We install metal roofs across Manatee County")
    r = admin.post("/api/ai/knowledge-bases/%s/files" % kb["id"], json={
        "filename": "services.pdf", "content_type": "application/pdf",
        "data": base64.b64encode(pdf).decode()})
    assert r.status_code == 201, r.text
    assert "metal roofs" in r.json()["body"] and r.json()["size_bytes"] == len(pdf)
    doc = _docx_bytes("Emergency tarping is available 24 hours a day.")
    r = admin.post("/api/ai/knowledge-bases/%s/files" % kb["id"], json={
        "filename": "policies.docx", "data": base64.b64encode(doc).decode()})
    assert r.status_code == 201, r.text
    assert "Emergency tarping" in r.json()["body"] and "25 years" in r.json()["body"]
    # The bytes are not kept anywhere: only the text, name, type and size.
    item = SessionLocal().scalar(select(AiKbItem).where(AiKbItem.title == "policies.docx"))
    assert item.content_type.endswith("wordprocessingml.document")
    assert not hasattr(item, "data")
    found = admin.post("/api/ai/knowledge-bases/search", json={
        "query": "do you do emergency tarps?", "knowledge_base_ids": [kb["id"]]}).json()
    assert found["useful"] and found["results"][0]["title"] == "policies.docx"
    found = admin.post("/api/ai/knowledge-bases/search", json={
        "query": "metal roof", "knowledge_base_ids": [kb["id"]]}).json()
    assert found["results"][0]["title"] == "services.pdf"


def test_an_unreadable_or_other_file_is_refused_and_writes_nothing(world):
    admin = world.client("admin")
    kb = admin.post("/api/ai/knowledge-bases", json={"name": "K"}).json()
    for filename, data in (("notes.txt", b"hello"), ("broken.pdf", b"%PDF-1.4 garbage"),
                           ("broken.docx", b"not a zip")):
        r = admin.post("/api/ai/knowledge-bases/%s/files" % kb["id"], json={
            "filename": filename, "data": base64.b64encode(data).decode()})
        assert r.status_code == 400, (filename, r.text)
    assert count(AiKbItem) == 0 and count(AiKbChunk) == 0


def test_retrieval_is_scoped_to_the_agents_attached_knowledge_bases(world, script):
    admin = world.client("admin")
    mine = admin.post("/api/ai/knowledge-bases", json={"name": "Retail"}).json()
    other = admin.post("/api/ai/knowledge-bases", json={"name": "Secret pricing"}).json()
    admin.post("/api/ai/knowledge-bases/%s/items" % mine["id"], json={
        "kind": "faq", "title": "Do you offer financing?",
        "body": "Yes, financing is available through our partner."})
    admin.post("/api/ai/knowledge-bases/%s/items" % other["id"], json={
        "kind": "article", "title": "Discount ladder",
        "body": "Internal: the minimum discount price for shingle replacement is 12 percent."})
    aid = make_agent(world, actions=["get_context", "search_knowledge", "report_knowledge_gap"],
                     knowledge_base_ids=[mine["id"]])
    script.responses.extend([
        anthropic_message([tool_use("search_knowledge", {"query": "discount price shingle"}, "a"),
                           tool_use("search_knowledge", {"query": "financing"}, "b")],
                          "tool_use"),
        anthropic_message([text("Financing yes; I'll check discounts.")]),
    ])
    world.client("dispatcher").post("/api/ai/agents/%s/run" % aid,
                                    json={"contact_id": world.ids["jane"]})
    drain()
    results = script.requests[1]["body"]["messages"][-1]["content"]
    assert "Discount ladder" not in results[0]["content"]
    assert "12 percent" not in str(script.requests)
    assert '"results": []' in results[0]["content"]
    assert "financing is available" in results[1]["content"]
    # The miss was recorded as a gap, with the run that saw it.
    gap = SessionLocal().scalar(select(AiKnowledgeGap))
    assert gap.question == "discount price shingle" and gap.count == 1 and gap.agent_id == aid


def test_gaps_are_deduplicated_and_resolving_one_creates_the_faq(world):
    db = SessionLocal()
    for q in ("Do you install solar panels?", "do you install SOLAR panels", "Other thing"):
        knowledge.record_gap(db, q, agent_id=1, run_id=7)
    db.commit()
    db.close()
    admin = world.client("admin")
    gaps = admin.get("/api/ai/knowledge-gaps").json()
    solar = next(g for g in gaps if "solar" in g["question"].lower())
    assert len(gaps) == 2 and solar["count"] == 2 and solar["last_run_id"] == 7
    kb = admin.post("/api/ai/knowledge-bases", json={"name": "FAQ"}).json()
    r = admin.post("/api/ai/knowledge-gaps/%s/resolve" % solar["id"], json={
        "knowledge_base_id": kb["id"], "answer": "No — roofing only."})
    assert r.status_code == 200 and r.json()["status"] == "resolved"
    faq = SessionLocal().scalar(select(AiKbItem))
    assert (faq.kind, faq.title, faq.body) == ("faq", "Do you install solar panels?",
                                               "No — roofing only.")
    hits = knowledge.RETRIEVER.search(SessionLocal(), [kb["id"]], "solar panels")
    assert hits and hits[0].item_id == faq.id
    assert admin.post("/api/ai/knowledge-gaps/%s/resolve" % solar["id"], json={
        "knowledge_base_id": kb["id"], "answer": "again"}).status_code == 409
    assert count(AiKbItem) == 1
    other = next(g for g in gaps if g["id"] != solar["id"])
    assert admin.post("/api/ai/knowledge-gaps/%s/dismiss" % other["id"]).json()["status"] == \
        "dismissed"
    assert admin.get("/api/ai/knowledge-gaps").json() == []
    # The dispatcher may read knowledge bases but not the gap list (it is built from logs).
    assert world.client("dispatcher").get("/api/ai/knowledge-gaps").status_code == 403
    assert world.client("dispatcher").get("/api/ai/knowledge-bases").status_code == 200


def test_a_knowledge_base_in_use_cannot_be_deleted(world):
    admin = world.client("admin")
    kb = admin.post("/api/ai/knowledge-bases", json={"name": "Retail"}).json()
    make_agent(world, mode="off", publish=False, knowledge_base_ids=[kb["id"]])
    r = admin.delete("/api/ai/knowledge-bases/%s" % kb["id"])
    assert r.status_code == 409 and "Follow-up" in r.json()["detail"]


def test_chunking_splits_long_text_and_keeps_every_word():
    long = " ".join("sentence number %d about roofing." % i for i in range(400))
    pieces = knowledge.chunk_text(long)
    assert len(pieces) > 3 and all(len(p) <= knowledge.CHUNK_CHARS for p in pieces)
    assert set(long.split()) <= set(" ".join(pieces).split())
    assert knowledge.terms("Roofs are leaking!") == ["roof", "leak"]
