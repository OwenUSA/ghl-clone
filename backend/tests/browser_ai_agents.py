"""Drive the AI Agents module in a REAL headless browser.

    cd backend && uv run python -m tests.browser_ai_agents [--keep] [--shots DIR]

Not collected by pytest (no `test_` prefix): it starts its own API and Vite dev server.
Needs node on PATH, the `capture` dependency group (Playwright) and its Chromium.

NOTHING HERE CAN REACH A MODEL PROVIDER OR A PHONE SYSTEM, by construction:

  * the API runs in a child process (`--serve`) in which `app.ai.providers.HTTP_CLIENT_FACTORY`
    returns an httpx2 client on a MockTransport that answers Anthropic Messages API JSON from
    a script and RAISES for any other URL; `crmlink`'s httpx calls raise too, and every
    owen-main URL is unset;
  * the dev server is `frontend/e2e/vite.sipmock.config.ts`: the SIP user is fake;
  * the database is a throwaway SQLite file in a fresh temp directory, made with
    `create_all` — never `app.seed`, never a database this did not create.

Every check asserts behaviour on screen AND behind it: what the fake provider was asked,
and what is in the database (the Try-it "Would" move must leave the deal where it was).
"""
from __future__ import annotations

import argparse
import contextlib
import io
import json
import os
import signal
import sqlite3
import subprocess
import sys
import tempfile
import traceback
from pathlib import Path

from tests.browser_dialer import Checks, free_port, wait_http

BACKEND = Path(__file__).resolve().parents[1]
FRONTEND = BACKEND.parent / "frontend"
ADMIN = "owner@ai.test"
PASSWORD = "ai-agents-check-password-1"
SECRET_KEY = "sk-ant-browser-check-9f8e7d6c"
KB_PHRASE = "shingle warranty lasts twenty five years"


# --------------------------------------------------------------------------- the API

def serve(port: int, db_path: str, record: str) -> None:
    """The API with the provider and the phone system replaced. Runs in a child process."""
    from cryptography.fernet import Fernet
    os.environ.update({
        "DATABASE_URL": "sqlite:///" + db_path,
        "AUTO_CREATE_ALL": "1",
        "AI_SECRETS_KEY": Fernet.generate_key().decode(),
    })
    for k in ("CRM_LINK_BASE_URL", "CRM_LINK_API_KEY", "OWEN_BASE_URL", "OWEN_SOFTPHONE_KEY",
              "ANTHROPIC_API_KEY", "OPENAI_API_KEY", "ANTHROPIC_BASE_URL", "OPENAI_BASE_URL"):
        os.environ.pop(k, None)
    sys.path.insert(0, str(BACKEND))
    import httpx2
    import uvicorn
    from app import auth, crmlink
    from app.ai import providers
    from app.db import SessionLocal
    from app.main import app
    from app.models import Contact, Opportunity, Pipeline, Role, Stage, User

    def forbidden(*a, **k):
        raise RuntimeError("a real HTTP request to the phone system was attempted")

    crmlink.httpx.post = forbidden
    crmlink.httpx.get = forbidden

    with SessionLocal() as db:
        db.add(User(email=ADMIN, name="Owner", role=Role.ADMIN,
                    password_hash=auth.hash_password(PASSWORD)))
        c = Contact(first_name="Jane", last_name="Doe", phone="+19415550101")
        db.add(c)
        p = Pipeline(name="Dream Team Roofing AHS")
        db.add(p)
        db.flush()
        s1 = Stage(pipeline_id=p.id, name="New Lead", position=0)
        s2 = Stage(pipeline_id=p.id, name="Inspection", position=1)
        db.add_all([s1, s2])
        db.flush()
        db.add(Opportunity(title="Jane Doe - roof leak", contact_id=c.id, pipeline_id=p.id,
                           stage_id=s1.id, value_cents=950000))
        db.commit()
        inspection = s2.id

    def log(entry: dict) -> None:
        with open(record, "a") as fh:
            fh.write(json.dumps(entry) + "\n")

    def message(content: list, stop: str, n_in: int = 120, n_out: int = 30):
        return httpx2.Response(200, json={
            "id": "msg_fake", "type": "message", "role": "assistant", "model": "claude-sonnet-5",
            "content": content, "stop_reason": stop,
            "usage": {"input_tokens": n_in, "output_tokens": n_out,
                      "cache_creation_input_tokens": 0, "cache_read_input_tokens": 0}})

    def handler(request):
        url = str(request.url)
        if not url.startswith("https://api.anthropic.com/"):
            log({"refused": url})
            raise RuntimeError("the fake provider only answers api.anthropic.com: " + url)
        log({"url": url, "method": request.method,
             "key_sent": request.headers.get("x-api-key") == SECRET_KEY,
             "body": json.loads(request.content) if request.content else None})
        if request.url.path == "/v1/models":
            return httpx2.Response(200, json={"data": [
                {"type": "model", "id": "claude-sonnet-5", "display_name": "Claude Sonnet 5",
                 "created_at": "2026-01-01T00:00:00Z"}], "has_more": False,
                "first_id": "claude-sonnet-5", "last_id": "claude-sonnet-5"})
        body = json.loads(request.content)
        if not body.get("tools"):
            return message([{"type": "text", "text": "OK"}], "end_turn", 12, 2)
        results = [b for m in body["messages"]
                   if m["role"] == "user" and isinstance(m["content"], list)
                   for b in m["content"] if b.get("type") == "tool_result"]
        if not results:
            return message([{"type": "text", "text": "Let me look at your job."},
                            {"type": "tool_use", "id": "toolu_1", "name": "get_context",
                             "input": {}}], "tool_use")
        if len(results) == 1:
            return message([{"type": "tool_use", "id": "toolu_2", "name": "move_stage",
                             "input": {"stage_id": inspection}}], "tool_use", 400, 20)
        return message([{"type": "text", "text": "Thanks Jane — I've asked the team to book "
                         "your inspection."}], "end_turn", 520, 25)

    providers.HTTP_CLIENT_FACTORY = lambda: httpx2.Client(
        transport=httpx2.MockTransport(handler))
    uvicorn.run(app, host="127.0.0.1", port=port, log_level="warning")


# --------------------------------------------------------------------------- fixtures

def make_pdf(text: str) -> bytes:
    """A one-page PDF with real extractable text, written by hand (offsets computed)."""
    stream = ("BT /F1 12 Tf 72 720 Td (%s) Tj ET" % text).encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
         b"/Resources << /Font << /F1 5 0 R >> >> >>"),
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
              % (len(objects) + 1, xref))
    return out.getvalue()


def make_docx(text: str) -> bytes:
    import docx
    d = docx.Document()
    d.add_heading("Storm policy", level=1)
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- the run

def run(shots: Path, keep: bool) -> int:
    from playwright.sync_api import expect, sync_playwright

    work = Path(tempfile.mkdtemp(prefix="ai-agents-check-"))
    db_path = str(work / "ai.db")
    record = work / "provider.jsonl"
    record.touch()
    pdf, docx_file = work / "warranty.pdf", work / "storm-policy.docx"
    pdf.write_bytes(make_pdf("Our " + KB_PHRASE + " from the install date."))
    docx_file.write_bytes(make_docx("After a storm we tarp an active leak within 24 hours."))
    api_port, vite_port = free_port(), free_port()
    procs: list[subprocess.Popen] = []
    checks = Checks()

    def q(sql: str):
        with sqlite3.connect(db_path) as c:
            return c.execute(sql).fetchall()

    def asked() -> list[dict]:
        return [json.loads(x) for x in record.read_text().splitlines() if x.strip()]

    try:
        procs.append(subprocess.Popen(
            [sys.executable, "-m", "tests.browser_ai_agents", "--serve", str(api_port), db_path,
             str(record)], cwd=BACKEND, start_new_session=True))
        wait_http(f"http://127.0.0.1:{api_port}/api/health")
        env = {**os.environ, "GHL_API_PORT": str(api_port), "GHL_VITE_PORT": str(vite_port)}
        procs.append(subprocess.Popen(
            ["npx", "vite", "--config", "e2e/vite.sipmock.config.ts"], cwd=FRONTEND, env=env,
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True))
        app_url = f"http://localhost:{vite_port}/"
        wait_http(app_url)
        print(f"API :{api_port} (pid {procs[0].pid}), Vite :{vite_port} (pid {procs[1].pid}),"
              f" db {db_path}", flush=True)
        stage_before = q("select stage_id from opportunities")[0][0]

        with sync_playwright() as pw:
            browser = pw.chromium.launch(headless=True)
            ctx = browser.new_context(viewport={"width": 1440, "height": 900})
            page = ctx.new_page()
            errors: list[str] = []
            page.on("pageerror", lambda e: errors.append(str(e)))
            page.goto(app_url)
            page.fill('input[type="email"]', ADMIN)
            page.fill('input[type="password"]', PASSWORD)
            page.click('button[type="submit"]')
            page.wait_for_selector("text=Sign in to continue", state="detached")

            # ---- the sidebar row and the bell ---------------------------------------------
            nav = page.get_by_role("button", name="AI Agents", exact=True)
            # The shell draws after the session query answers; under load that is not instant.
            with contextlib.suppress(Exception):
                nav.first.wait_for(timeout=20000)
            checks.ok(nav.count() == 1, "AI Agents is a real sidebar row for an admin")
            checks.ok(page.get_by_role("button", name="Alerts").count() == 1,
                      "the alert bell is on the page beside the status dot")

            # ---- Settings → AI Connections ------------------------------------------------
            page.get_by_role("button", name="Settings", exact=True).click()
            page.get_by_role("tab", name="AI Connections").click()
            expect(page.get_by_text("Pause all AI agents", exact=True)).to_be_visible()
            page.get_by_role("button", name="Add connection").first.click()
            dlg = page.get_by_role("dialog", name="Add connection")
            dlg.get_by_label("Connection name").fill("Claude — office")
            checks.ok(dlg.get_by_label("Default model").input_value() == "claude-sonnet-5",
                      "a new Anthropic connection is prefilled with claude-sonnet-5")
            dlg.get_by_label("API key").fill(SECRET_KEY)
            dlg.get_by_role("button", name="Test", exact=True).click()
            status = dlg.locator("[data-test-ok]")
            expect(status).to_be_visible(timeout=10000)
            checks.ok(status.get_attribute("data-test-ok") == "true"
                      and "Connected" in status.inner_text(),
                      f"Test shows OK with latency ({status.inner_text()!r})")
            checks.ok(q("select count(*) from ai_connections")[0][0] == 0,
                      "Test wrote nothing — no connection saved yet")
            page.screenshot(path=str(shots / "01-connection-modal.png"))
            dlg.get_by_role("button", name="Add connection").click()
            expect(page.get_by_text("•••• 7d6c")).to_be_visible(timeout=5000)
            stored = q("select api_key_encrypted from ai_connections")[0][0]
            checks.ok(SECRET_KEY not in stored and SECRET_KEY not in page.content(),
                      "the key is encrypted at rest and never back on the page")
            page.screenshot(path=str(shots / "02-ai-connections.png"))

            # ---- AI Agents → Create agent → builder ---------------------------------------
            nav.click()
            expect(page.get_by_role("tab", name="Agent Logs")).to_be_visible()
            page.screenshot(path=str(shots / "03-agents-empty.png"))
            page.get_by_role("button", name="Create agent").click()
            cdlg = page.get_by_role("dialog", name="Create agent")
            checks.ok(cdlg.locator("option[value=voice]").count() == 0,
                      "Voice is not offered as a channel")
            cdlg.get_by_label("Agent name").fill("Lead follow-up")
            cdlg.get_by_role("button", name="Create", exact=True).click()
            expect(page.get_by_role("navigation", name="Agent sections")).to_be_visible(
                timeout=5000)
            checks.ok(q("select mode from ai_agents")[0][0] == "off", "the new agent is Off")

            sections = page.get_by_role("navigation", name="Agent sections")
            sections.get_by_role("button", name="Persona & goals").click()
            page.get_by_label("Role / persona").fill("You are the friendly office assistant.")
            page.get_by_role("button", name="Add", exact=True).click()
            page.get_by_label("Goals 1", exact=True).fill("Get the roof inspected")
            sections.get_by_role("button", name="AI connection").click()
            page.get_by_label("Connection", exact=True).select_option(label="Claude — office")
            checks.ok(page.get_by_label("Model", exact=True).input_value() == "claude-sonnet-5",
                      "choosing a connection fills its default model")
            sections.get_by_role("button", name="Actions").click()
            page.locator("label", has_text="Move stage").locator("input[type=checkbox]").check()
            checks.ok(page.get_by_text("Transfer call").count() == 0,
                      "voice-only actions are not offered to a Text agent")
            sections.get_by_role("button", name="Triggers").click()
            page.get_by_label("Trigger to add").select_option("manual")
            page.get_by_role("button", name="Add trigger").click()
            page.get_by_role("button", name="Save draft").click()
            expect(page.get_by_text("Draft saved.")).to_be_visible(timeout=5000)
            draft = json.loads(q("select draft from ai_agents")[0][0])
            checks.ok(draft["connection_id"] == 1 and "move_stage" in draft["actions"]
                      and draft["goals"] == ["Get the roof inspected"]
                      and draft["triggers"] == [{"type": "manual"}],
                      "Save draft stored the sections")
            checks.ok(q("select count(*) from ai_agent_versions")[0][0] == 0,
                      "saving a draft published nothing")
            sections.get_by_role("button", name="Compiled prompt").click()
            expect(page.get_by_label("Compiled prompt")).to_contain_text("Get the roof inspected",
                                                                         timeout=5000)
            checks.ok(True, "the compiled prompt preview shows the goal")

            # ---- Try it, as the seeded opportunity ----------------------------------------
            tri = page.get_by_role("complementary", name="Try it")
            tri.get_by_label("Test as contact").fill("Jane")
            tri.get_by_role("option").first.click()
            tri.get_by_label("Test opportunity").select_option(label="Jane Doe - roof leak")
            tri.get_by_label("Message as the customer").fill(
                "Hi, my roof is leaking, can someone come?")
            tri.get_by_role("button", name="Send").click()
            would = tri.locator("[data-would]")
            expect(would.first).to_be_visible(timeout=15000)
            card = would.first.inner_text()
            checks.ok("Would: Move" in card and "Inspection" in card,
                      f"a Would card shows the move ({card!r})")
            reply = tri.locator("[data-role=assistant]").last.inner_text()
            checks.ok(reply.startswith("Thanks Jane"), "the agent's reply appears")
            checks.ok(q("select stage_id from opportunities")[0][0] == stage_before,
                      "Try-it moved nothing: the opportunity is still in its stage")
            calls = [a for a in asked() if a.get("body") and a["body"].get("tools")]
            checks.ok(len(calls) == 3 and all(a["key_sent"] for a in calls),
                      "the provider was called three times with the stored key")
            checks.ok(not any(SECRET_KEY in json.dumps(a["body"]) for a in calls),
                      "the key is never inside a request body")
            checks.ok(not any("refused" in a for a in asked()), "no request went anywhere else")
            checks.ok(q("select is_test, outcome from ai_runs") == [(1, "completed")],
                      "the test run is logged as a completed test")
            page.screenshot(path=str(shots / "04-builder-try-it.png"))

            # ---- Publish -------------------------------------------------------------------
            page.get_by_role("button", name="Publish", exact=True).click()
            expect(page.get_by_text("Published v1")).to_be_visible(timeout=5000)
            checks.ok(q("select version from ai_agent_versions") == [(1,)],
                      "Publish froze version 1")
            page.get_by_role("radio", name="Auto-pilot").click()
            confirm = page.get_by_role("alertdialog", name="Switch to Auto-pilot?")
            checks.ok(confirm.is_visible(), "Auto-pilot asks for confirmation")
            confirm.get_by_role("button", name="Cancel").click()
            checks.ok(q("select mode from ai_agents")[0][0] == "off",
                      "cancelling the confirmation changed nothing")

            # ---- Knowledge Base -----------------------------------------------------------
            page.get_by_role("tab", name="Knowledge Base").click()
            page.get_by_role("button", name="New knowledge base").click()
            kdlg = page.get_by_role("dialog", name="New knowledge base")
            kdlg.get_by_label("Knowledge base name").fill("Roofing basics")
            kdlg.get_by_role("button", name="Create").click()
            page.get_by_role("tab", name="Files (0)").click()
            page.get_by_label("Upload file").set_input_files(str(pdf))
            expect(page.get_by_text("characters of text")).to_be_visible(timeout=10000)
            page.get_by_label("Upload file").set_input_files(str(docx_file))
            expect(page.get_by_role("tab", name="Files (2)")).to_be_visible(timeout=10000)
            rows = page.locator("table tbody tr")
            checks.ok(rows.count() == 2 and "warranty.pdf" in rows.all_inner_texts()[1]
                      + rows.all_inner_texts()[0],
                      "both files are listed")
            files = q("select title, content_type, size_bytes, length(body) from ai_kb_items "
                      "order by id")
            checks.ok(len(files) == 2 and files[0][1] == "application/pdf" and files[0][3] > 20
                      and files[1][0] == "storm-policy.docx" and files[1][3] > 20,
                      f"PDF and .docx text was extracted and stored ({files})")
            checks.ok(page.locator("table").inner_text().count("Characters") == 1,
                      "the Files table shows a Characters column")
            page.get_by_label("Test search").fill("how long is the shingle warranty")
            page.locator("form").get_by_role("button", name="Search", exact=True).click()
            expect(page.get_by_label("Search results")).to_contain_text("warranty.pdf",
                                                                        timeout=5000)
            checks.ok(KB_PHRASE.split()[0] in page.get_by_label("Search results").inner_text(),
                      "Test search finds the PDF's text")
            page.screenshot(path=str(shots / "05-knowledge-base.png"))

            # ---- Agent Logs ---------------------------------------------------------------
            page.get_by_role("tab", name="Agent Logs").click()
            row = page.locator("[data-run]").first
            expect(row).to_be_visible(timeout=5000)
            checks.ok("Test" in row.inner_text() and "Completed" in row.inner_text(),
                      "the test run is listed with its outcome")
            page.screenshot(path=str(shots / "06-agent-logs.png"))
            row.click()
            rdlg = page.get_by_role("dialog", name="Run #1")
            expect(rdlg.locator("[data-step=tool_call]").first).to_be_visible(timeout=5000)
            checks.ok(rdlg.locator("[data-step=tool_call]").count() == 2,
                      "the transcript shows both tool calls")
            checks.ok(rdlg.locator("[data-step=action][data-status=would]").count() == 1,
                      "the transcript shows the Would action")
            page.screenshot(path=str(shots / "07-run-detail.png"))
            rdlg.get_by_role("button", name="Close").click()
            page.get_by_role("tab", name="Metrics").click()
            expect(page.locator("[data-metric=runs]")).to_be_visible(timeout=5000)
            checks.ok(page.locator("[data-metric=runs]").inner_text() == "0",
                      "Metrics renders and does not count the Try-it test")
            page.screenshot(path=str(shots / "08-metrics.png"))

            checks.ok(not errors, f"no uncaught page errors ({errors[:2]})")
            browser.close()
    except Exception:  # noqa: BLE001 - any failure is one failed check, then clean up
        traceback.print_exc()
        checks.ok(False, "the run completed without an exception")
    finally:
        for p in procs:
            with contextlib.suppress(ProcessLookupError):
                os.killpg(p.pid, signal.SIGTERM)
        for p in procs:
            try:
                p.wait(timeout=10)
            except subprocess.TimeoutExpired:
                os.killpg(p.pid, signal.SIGKILL)
        if not keep:
            for f in work.iterdir():
                f.unlink()
            work.rmdir()

    passed = sum(1 for ok, _ in checks.results if ok)
    print(f"\n{passed}/{len(checks.results)} AI Agents browser checks passed")
    return 0 if passed == len(checks.results) else 1


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--serve":
        serve(int(sys.argv[2]), sys.argv[3], sys.argv[4])
        sys.exit(0)
    ap = argparse.ArgumentParser()
    ap.add_argument("--keep", action="store_true", help="keep the temp database")
    ap.add_argument("--shots", default=None, help="where to write screenshots")
    args = ap.parse_args()
    out = Path(args.shots or tempfile.mkdtemp(prefix="ai-agents-shots-"))
    out.mkdir(parents=True, exist_ok=True)
    print("screenshots:", out)
    sys.exit(run(out, args.keep))
