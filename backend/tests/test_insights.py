"""Daily report & insights (app/insights), by behaviour - the important ones.

Tagger (provider scripted): "still leaking after the repair" becomes ONE tag with its quote,
linked to the job and its technician; an invented quote is dropped; nothing is read twice; an
excluded phone is never sent to the model.
Metrics: a callback is charged to the ORIGINAL job's technician; still leaking counts distinct
jobs; upsell money; a hidden board is left out; week / month / quarter boundaries.
Alerts: still leaking with no future visit -> one Dispatch item + one bell; a booked visit
resolves it.
Ask (scripted): a breakdown tool call returns a table + bar chart; a number the tools never
returned is flagged. API: a technician 403; pin -> the report shows it; the Excel has charts.
"""
import io
import json
from datetime import UTC, date, datetime, timedelta

import pytest
from app.ai import vault
from app.dispatch import ai as dai
from app.insights import ask, config, metrics, service, tagger
from app.models import (
    AgentRegistryEntry,
    AiAlert,
    AiConnection,
    Contact,
    Conversation,
    ConversationEvent,
    ConversationRead,
    ConversationSignal,
    Direction,
    DispatchItem,
    DispatchJob,
    EventType,
    Role,
    User,
    ZuperRecordVersion,
    ZuperStatusHistory,
)
from openpyxl import load_workbook
from tests.ai_support import openai_completion

NOW = datetime(2026, 10, 22, 16, 0, tzinfo=UTC)               # Thu noon New York
LEAK = "Hi this is Ana, the roof is still leaking after the repair you did on Monday, please call"


@pytest.fixture()
def ds(db, secrets_key):
    conn = AiConnection(name="Deepseek", provider="deepseek",
                        api_key_encrypted=vault.encrypt("sk-2"), api_key_last4="2",
                        default_model="deepseek-flash")
    db.add(conn)
    db.flush()
    s = dai.settings(db, for_update=True)
    s.ai_enabled, s.connection_id, s.model = True, conn.id, "deepseek-flash"
    db.commit()


def _job(db, uid, number, tech, *, board="Retail", phones=None, start=None):
    db.add(DispatchJob(job_uid=uid, job_number=number, board=board, status="Repair Complete",
                       customer_name="Ana Smith", phones=phones or [], is_open=True,
                       technician=tech, scheduled_start=start))


def _text(db, phone, body, at):
    c = Contact(first_name="Ana", last_name="Smith", phone="+1" + phone)
    db.add(c)
    db.flush()
    conv = Conversation(contact_id=c.id)
    db.add(conv)
    db.flush()
    e = ConversationEvent(conversation_id=conv.id, type=EventType.SMS,
                          direction=Direction.INBOUND, body=body, occurred_at=at)
    db.add(e)
    db.commit()
    return e


def _answer(tags: dict) -> dict:
    return openai_completion(content=json.dumps(tags), model="deepseek-flash")


# ------------------------------------------------------------------------- tagger

def test_a_still_leaking_text_is_tagged_once_with_its_quote(db, ds, script):
    _job(db, "job-1", "701", "Luis", phones=["9415550101"])
    _text(db, "9415550101", LEAK, NOW - timedelta(hours=2))
    _text(db, "9415550202", "Can you give me a quote for gutters on the back of my house",
          NOW - timedelta(hours=1))
    script.responses = [_answer({
        "0": [{"tag": "still_leaking", "quote": "the roof is still leaking after the repair"}],
        "1": [{"tag": "upsell_interest", "quote": "a brand new roof for the whole house"}]})]
    out = tagger.tag(db, NOW, config.Settings())
    sig = db.query(ConversationSignal).one()                      # the invented quote dropped
    assert (sig.tag, sig.job_number, sig.technician) == ("still_leaking", "701", "Luis")
    assert out == {"read": 2, "tags": 1} and db.query(ConversationRead).count() == 2
    assert tagger.tag(db, NOW, config.Settings()) == {"read": 0, "tags": 0}
    assert script.calls == 1                                       # nothing read twice


def test_an_excluded_phone_never_reaches_the_model(db, ds, script):
    _text(db, "9545550000", LEAK, NOW - timedelta(hours=1))
    assert tagger.tag(db, NOW, config.Settings(exclude_phones=["9545550000"]))["read"] == 0
    assert script.calls == 0


# ------------------------------------------------------------------------- metrics

def _version(db, uid, record):
    db.add(ZuperRecordVersion(module="job", zuper_uid=uid, content_hash=uid,
                              record=dict(record, job_uid=uid)))


def test_metrics_callbacks_leaks_upsells_and_hidden_boards(db):
    week = metrics.period("week", NOW.date())
    created = "2026-10-20T15:00:00Z"
    tech = [{"label": "Technician", "value": "Luis"}]
    _version(db, "orig", {"work_order_number": "500", "custom_fields": tech,
                          "job_category": {"category_name": "Retail"},
                          "scheduled_start_time": "2026-10-10T14:00:00Z",
                          "created_at": "2026-10-01T14:00:00Z"})
    _version(db, "cb", {"work_order_number": "510", "job_type": "REVISIT",
                        "parent_job": {"job_uid": "orig"},
                        "custom_fields": [{"label": "Technician", "value": "Pedro"}],
                        "job_category": {"category_name": "Retail"}, "created_at": created})
    _version(db, "ahs", {"work_order_number": "600", "job_type": "REVISIT",
                         "job_category": {"category_name": "AHS"}, "created_at": created})
    db.add(ZuperStatusHistory(history_uid="h1", job_uid="orig", job_number="500",
                              category_name="Retail", status_name="Repair Complete",
                              changed_at=NOW - timedelta(days=1), checklist=[
                                  {"question": "Customer chose", "answer": "Repair + gutters"},
                                  {"question": "Customer paid ($)", "answer": "1,200"}]))
    for i, job in enumerate(("job-1", "job-1", "job-2")):
        db.add(ConversationSignal(source="quo", source_ref="ce:%d" % i, kind="text",
                                  at=NOW - timedelta(hours=i + 1), phone="9415550101",
                                  inbound=True, tag="still_leaking", quote="still leaking",
                                  job_uid=job, board="Retail", technician="Luis"))
    db.commit()
    rep = metrics.compute(db, *week)
    t = rep["tiles"]
    assert (t["callbacks"], t["still_leaking"], t["upsells"], t["upsell_paid"]) == (
        2, 2, 1, 1200.0)
    luis = next(r for r in rep["tables"]["technicians"] if r["technician"] == "Luis")
    assert (luis["callbacks"], luis["repairs"], luis["callback_rate"], luis["still_leaking"]
            ) == (1, 1, 1.0, 2)                                    # charged to the original's
    assert rep["tables"]["callbacks"][0]["days_later"] is not None
    hidden = metrics.compute(db, *week, hidden={"AHS"})["tiles"]
    assert hidden["callbacks"] == 1


def test_period_boundaries():
    d = date(2026, 10, 22)
    assert metrics.period("week", d) == (date(2026, 10, 19), date(2026, 10, 25))
    assert metrics.period("month", d) == (date(2026, 10, 1), date(2026, 10, 31))
    assert metrics.period("quarter", d) == (date(2026, 10, 1), date(2026, 12, 31))
    assert metrics.previous(date(2026, 10, 1), date(2026, 12, 31), "quarter") == (
        date(2026, 7, 1), date(2026, 9, 30))
    assert metrics.previous(*metrics.period("week", d), "week")[0] == date(2026, 10, 12)


# ------------------------------------------------------------------------- alerts

def test_still_leaking_raises_one_issue_and_a_booked_visit_resolves_it(db, monkeypatch):
    monkeypatch.setenv("INSIGHTS_ENABLED", "true")
    db.add(User(email="boss@x.test", name="Boss", role=Role.ADMIN))
    db.add(AgentRegistryEntry(key="insights", config={}))
    _job(db, "job-1", "701", "Luis", phones=["9415550101"])
    db.add(ConversationSignal(source="quo", source_ref="ce:1", kind="text",
                              at=NOW - timedelta(hours=1), phone="9415550101", inbound=True,
                              tag="still_leaking", quote="still leaking after the repair",
                              job_uid="job-1", job_number="701", board="Retail",
                              technician="Luis"))
    db.add(ConversationSignal(source="quo", source_ref="ce:2", kind="text",
                              at=NOW - timedelta(hours=1), phone="9415550101", inbound=True,
                              tag="praise", quote="thank you so much"))
    db.commit()
    out = {"alerts": 0, "resolved": 0}
    service.raise_alerts(db, NOW, out)
    service.raise_alerts(db, NOW + timedelta(hours=1), out)
    db.commit()
    item = db.query(DispatchItem).one()
    assert (item.kind, item.queue, item.urgent, out["alerts"]) == (
        "customer_issue", "issues", True, 1)
    assert "STILL LEAKING" in item.title and "Luis" in item.why
    assert db.query(AiAlert).count() == 1
    db.query(DispatchJob).one().scheduled_start = NOW + timedelta(days=2)
    db.commit()
    service.raise_alerts(db, NOW + timedelta(hours=2), out)
    db.commit()
    assert db.query(DispatchItem).one().state == "resolved"


def test_the_pass_reads_nothing_while_the_server_says_off(db, ds, script):
    _text(db, "9415550101", LEAK, NOW - timedelta(hours=1))
    assert "INSIGHTS_ENABLED" in service.run(db, NOW)["stopped"] and script.calls == 0


# ------------------------------------------------------------------------- ask + API

def _call(name, **args):
    return openai_completion(tool_calls=[{"id": "c1", "type": "function", "function": {
        "name": name, "arguments": json.dumps(args)}}], finish="tool_calls",
        model="deepseek-flash")


def test_ask_returns_a_table_and_flags_a_number_it_never_saw(db, ds, script):
    db.add(ConversationSignal(source="quo", source_ref="ce:1", kind="text",
                              at=NOW - timedelta(hours=1), phone="9415550101", inbound=True,
                              tag="still_leaking", quote="still leaking",
                              job_uid="job-1", board="Retail", technician="Luis"))
    db.commit()
    script.responses = [_call("breakdown", table="technicians", range="week",
                              date="2026-10-22"),
                        openai_completion(content="Luis has 1 still leaking; 47 callbacks.",
                                          model="deepseek-flash")]
    res = ask.answer(db, "still leaking per technician this week", hidden=set(), now=NOW)
    assert res["result"]["chart"] == "bar" and res["result"]["rows"][0][0] == "Luis"
    assert res["spec"] == {"tool": "breakdown", "args": {"table": "technicians",
                                                         "range": "week", "date": "2026-10-22"}}
    assert "check them: 47" in res["reply"]


def test_the_api_access_pin_and_excel(world):
    assert world.client("tech").get("/api/insights/report").status_code == 403
    d = world.client("dispatcher")
    r = d.post("/api/insights/saved", json={"name": "Leaks by tech", "spec": {
        "tool": "breakdown", "args": {"table": "technicians", "range": "month"}}})
    assert r.status_code == 200, r.text
    rep = d.get("/api/insights/report", params={"range": "month"}).json()
    assert [p["name"] for p in rep["pinned"]] == ["Leaks by tech"]
    assert {t["key"] for t in rep["tiles"]} >= {"callbacks", "still_leaking", "upsell_paid"}
    x = d.get("/api/insights/report.xlsx", params={"range": "month"})
    wb = load_workbook(io.BytesIO(x.content))
    assert "Summary" in wb.sheetnames and "Technicians" in wb.sheetnames
    csv = d.get("/api/insights/table/callbacks.csv").content.decode("utf-8-sig")
    assert csv.splitlines()[0]
    assert d.delete("/api/insights/saved/%d" % rep["pinned"][0]["id"]).json()["deleted"]
